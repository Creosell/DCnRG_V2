"""Converts manually filled "9 points" Excel sheets into DCnRG_V2 data/*.json files.

Sheet layout (per 9-row block, two samples side by side, columns offset by +6):
    row+0, col A/G       -> sample index
    row+0, col C/I        -> serial number (merged C1:F1 / I1:L1)
    row+2..+4, col C:E/I:K -> 3x3 brightness (Lv) grid: Top/Middle/Bottom x Left/Center/Right
    row+5, col D/J        -> color temperature (T)
    row+6, col C/I        -> White Lv
    row+6, col D/J        -> Black Lv
    row+7, col B:E/H:K    -> White/Red/Green/Blue x
    row+8, col B:E/H:K    -> White/Red/Green/Blue y
"""

import argparse
import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl

SHEET_NAME = "9 points"
BLOCK_ROWS = 9
SAMPLE_COL_OFFSETS = (0, 6)  # left sample starts at column A, right sample at column G


def _device_name_from_filename(xlsx_path: Path) -> str:
    match = re.match(r"(.+?)\s+ORDER\b", xlsx_path.stem, flags=re.IGNORECASE)
    return match.group(1) if match else xlsx_path.stem


def _extract_sample(ws, row0: int, col0: int):
    """Extracts one sample's raw cells starting at row0/col0 (1-indexed, base=A or G)."""
    idx = ws.cell(row=row0, column=col0).value
    if idx is None:
        return None

    serial = ws.cell(row=row0, column=col0 + 2).value

    top = [ws.cell(row=row0 + 2, column=col0 + c).value for c in (2, 3, 4)]
    mid = [ws.cell(row=row0 + 3, column=col0 + c).value for c in (2, 3, 4)]
    bot = [ws.cell(row=row0 + 4, column=col0 + c).value for c in (2, 3, 4)]

    temperature = ws.cell(row=row0 + 5, column=col0 + 3).value
    white_lv = ws.cell(row=row0 + 6, column=col0 + 2).value
    black_lv = ws.cell(row=row0 + 6, column=col0 + 3).value

    gamut_x = [ws.cell(row=row0 + 7, column=col0 + c).value for c in (1, 2, 3, 4)]
    gamut_y = [ws.cell(row=row0 + 8, column=col0 + c).value for c in (1, 2, 3, 4)]

    return {
        "idx": idx,
        "serial": serial,
        "top": top,
        "mid": mid,
        "bot": bot,
        "temperature": temperature,
        "white_lv": white_lv,
        "black_lv": black_lv,
        "white_xy": (gamut_x[0], gamut_y[0]),
        "red_xy": (gamut_x[1], gamut_y[1]),
        "green_xy": (gamut_x[2], gamut_y[2]),
        "blue_xy": (gamut_x[3], gamut_y[3]),
    }


def _build_report(sample: dict, device_config: str, is_tv: bool, measurement_dt: str) -> dict:
    top_left, top_center, top_right = sample["top"]
    mid_left, center, mid_right = sample["mid"]
    bot_left, bot_center, bot_right = sample["bot"]

    def point(location, lv, xy=(None, None), t=None):
        x, y = xy
        return {"Location": location, "x": x, "y": y, "Lv": lv, "T": t}

    measurements = [
        point("TopLeft", top_left),
        point("TopCenter", top_center),
        point("TopRight", top_right),
        point("MiddleLeft", mid_left),
        point("Center", center, t=sample["temperature"]),
        point("MiddleRight", mid_right),
        point("BottomLeft", bot_left),
        point("BottomCenter", bot_center),
        point("BottomRight", bot_right),
        point("WhiteColor", sample["white_lv"], xy=sample["white_xy"]),
        point("RedColor", None, xy=sample["red_xy"]),
        point("GreenColor", None, xy=sample["green_xy"]),
        point("BlueColor", None, xy=sample["blue_xy"]),
        point("BlackColor", sample["black_lv"]),
    ]

    return {
        "SerialNumber": str(sample["serial"]) if sample["serial"] is not None else str(sample["idx"]),
        "DeviceConfiguration": device_config,
        "IsTV": is_tv,
        "MeasurementDateTime": measurement_dt,
        "Measurements": measurements,
    }


def convert(xlsx_path: Path, output_dir: Path, is_tv: bool = True) -> list[Path]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET_NAME]
    device_config = _device_name_from_filename(xlsx_path)
    measurement_dt = datetime.now().strftime("%Y%m%d_%H%M%S")

    output_dir.mkdir(parents=True, exist_ok=True)
    written = []

    for row0 in range(1, ws.max_row + 1, BLOCK_ROWS):
        for col_offset in SAMPLE_COL_OFFSETS:
            sample = _extract_sample(ws, row0, 1 + col_offset)
            if sample is None or sample["serial"] is None:
                continue

            report = _build_report(sample, device_config, is_tv, measurement_dt)
            safe_serial = re.sub(r"[^A-Za-z0-9_.-]", "_", report["SerialNumber"])
            out_path = output_dir / f"{safe_serial}.json"
            out_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
            written.append(out_path)

    return written


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx_file", type=Path, help="Path to the manually filled Excel file")
    parser.add_argument("--output-dir", type=Path, default=Path("data"), help="Directory for generated JSON files")
    parser.add_argument("--monitor", action="store_true", help="Mark device as monitor instead of TV")
    args = parser.parse_args()

    written = convert(args.xlsx_file, args.output_dir, is_tv=not args.monitor)
    print(f"Written {len(written)} file(s) to {args.output_dir}")
    for path in written:
        print(f"  {path}")


if __name__ == "__main__":
    main()
